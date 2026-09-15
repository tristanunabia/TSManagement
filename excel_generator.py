import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
import datetime
import calendar
from date_engine import generate_cutoff_dates, get_previous_month

# Color constants
PRIMARY_COLOR = "1F4E78"    # Navy
SECONDARY_COLOR = "D9E1F2"  # Ice Blue
WEEKEND_COLOR = "F2F2F2"    # Light Grey
TEXT_LIGHT = "FFFFFF"       # White
BORDER_COLOR = "D9D9D9"     # Thin Grey
ACCENT_GREEN = "E2EFDA"     # Soft Green (LEC)
ACCENT_BLUE = "DDEBF7"      # Soft Blue (LAB)
ACCENT_ORANGE = "FFF2CC"    # Soft Orange (Consultation)

def get_thin_border():
    thin = Side(border_style="thin", color=BORDER_COLOR)
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def generate_excel_timesheet(instructor_row, class_loads_df, year, month, selected_cutoff, timesheet_data):
    """
    Generates an Excel workbook byte stream.
    instructor_row: dict or sqlite3.Row with instructor details
    class_loads_df: Pandas DataFrame of instructor's class loads
    year, month, selected_cutoff: parameters of the selected timesheet
    timesheet_data: dict containing edited cell values {(row_name, date_str): hours}
    """
    wb = openpyxl.Workbook()
    
    # Remove the default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)
    
    # 1. Create Sched sheet
    create_sched_sheet(wb, instructor_row, class_loads_df)
    
    # 2. Create Class Load sheet
    create_class_load_sheet(wb, class_loads_df)
    
    # 3. Create Timesheet sheets
    # We create BOTH sheet structures, but we hide the inactive one
    ws_11_25 = create_timesheet_sheet(wb, "Timesheet 11-25", instructor_row, class_loads_df, year, month, "11-25", timesheet_data)
    ws_26_10 = create_timesheet_sheet(wb, "Timesheet 26-10", instructor_row, class_loads_df, year, month, "26-10", timesheet_data)
    
    # Control sheet visibility
    if selected_cutoff == "11-25":
        ws_26_10.sheet_state = 'hidden'
        ws_11_25.sheet_state = 'visible'
        # Set active sheet to 11-25 (which is index 2, Sched=0, Class Load=1, 11-25=2)
        wb.active = wb.sheetnames.index("Timesheet 11-25")
    else:
        ws_11_25.sheet_state = 'hidden'
        ws_26_10.sheet_state = 'visible'
        wb.active = wb.sheetnames.index("Timesheet 26-10")
        
    # Write workbook to in-memory bytes
    fp = io.BytesIO()
    wb.save(fp)
    fp.seek(0)
    return fp.getvalue()

def create_sched_sheet(wb, instructor, class_loads):
    ws = wb.create_sheet(title="Sched")
    ws.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws.merge_cells("A1:H1")
    ws["A1"] = "WEEKLY MASTER SCHEDULE"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=TEXT_LIGHT)
    ws["A1"].fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40
    
    # Faculty Details
    details = [
        ("Faculty Name:", instructor["name"], "Dept/College:", instructor["department"]),
        ("Status:", instructor["employment_status"], "Email:", instructor.get("email", "N/A"))
    ]
    for r_idx, row_data in enumerate(details, start=2):
        ws.row_dimensions[r_idx].height = 20
        ws.cell(row=r_idx, column=1, value=row_data[0]).font = Font(name="Calibri", bold=True)
        ws.cell(row=r_idx, column=2, value=row_data[1])
        ws.cell(row=r_idx, column=4, value=row_data[2]).font = Font(name="Calibri", bold=True)
        ws.cell(row=r_idx, column=5, value=row_data[3])
    
    # Schedule Grid Headers
    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    ws.cell(row=5, column=1, value="Time").font = Font(bold=True, color=TEXT_LIGHT)
    ws.cell(row=5, column=1).fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    ws.cell(row=5, column=1).alignment = Alignment(horizontal="center")
    
    for c_idx, day in enumerate(days, start=2):
        cell = ws.cell(row=5, column=c_idx, value=day)
        cell.font = Font(bold=True, color=TEXT_LIGHT)
        cell.fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        
    ws.row_dimensions[5].height = 25
    
    # Generate Time Slots (06:30 to 21:00, 15 min steps)
    time_slots = []
    start_t = datetime.time(6, 30)
    end_t = datetime.time(21, 0)
    curr_t = datetime.datetime.combine(datetime.date.today(), start_t)
    end_dt = datetime.datetime.combine(datetime.date.today(), end_t)
    
    while curr_t <= end_dt:
        time_slots.append(curr_t.time())
        curr_t += datetime.timedelta(minutes=15)
        
    # Map weekdays to columns
    day_col_map = {"M": 2, "T": 3, "W": 4, "TH": 5, "F": 6, "S": 7, "SU": 8}
    
    # Initialize empty cells with thin borders
    for r_idx, t in enumerate(time_slots, start=6):
        ws.row_dimensions[r_idx].height = 18
        # Format time display
        t_str = t.strftime("%H:%M")
        cell = ws.cell(row=r_idx, column=1, value=t_str)
        cell.font = Font(size=9)
        cell.alignment = Alignment(horizontal="center")
        cell.border = get_thin_border()
        
        for col_idx in range(2, 9):
            grid_cell = ws.cell(row=r_idx, column=col_idx)
            grid_cell.border = get_thin_border()
            # Highlight Saturday (col 7) and Sunday (col 8)
            if col_idx in (7, 8):
                grid_cell.fill = PatternFill(start_color=WEEKEND_COLOR, end_color=WEEKEND_COLOR, fill_type="solid")
                
    # Place classes and merge slots
    # Standard time intervals are 15 mins. E.g. start at 08:30, end at 10:00 is 6 slots.
    # We find the start row and end row, then merge.
    for _, row in class_loads.iterrows():
        day = row["day_of_week"]
        col_idx = day_col_map.get(day)
        if not col_idx:
            continue
            
        try:
            start_str = row["start_time"]
            end_str = row["end_time"]
            
            s_time = datetime.datetime.strptime(start_str, "%H:%M").time()
            e_time = datetime.datetime.strptime(end_str, "%H:%M").time()
            
            # Find row index matching start_time and end_time
            start_r_idx = -1
            end_r_idx = -1
            
            for r_idx, t in enumerate(time_slots, start=6):
                if t == s_time:
                    start_r_idx = r_idx
                if t == e_time:
                    end_r_idx = r_idx - 1 # class finishes before this final slot starts
                    
            if start_r_idx != -1 and end_r_idx != -1 and start_r_idx <= end_r_idx:
                # Merge cells in this column from start_r_idx to end_r_idx
                ws.merge_cells(start_row=start_r_idx, start_column=col_idx, end_row=end_r_idx, end_column=col_idx)
                
                # Style top cell
                top_cell = ws.cell(row=start_r_idx, column=col_idx)
                top_cell.value = f"{row['subject_name']}\n{row['section_code']}\n{row['room']}\n({row['type']})"
                top_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                top_cell.font = Font(size=9, bold=True)
                
                # Determine soft color based on type & category
                cat = row["load_category"]
                ctype = row["type"]
                if cat == "Consultation":
                    fill_c = ACCENT_ORANGE
                elif ctype == "LEC":
                    fill_c = ACCENT_GREEN
                else:
                    fill_c = ACCENT_BLUE
                    
                p_fill = PatternFill(start_color=fill_c, end_color=fill_c, fill_type="solid")
                for r in range(start_r_idx, end_r_idx + 1):
                    ws.cell(row=r, column=col_idx).fill = p_fill
        except Exception as e:
            # Skip invalid times or formats
            continue
            
    # Set column widths
    ws.column_dimensions["A"].width = 10
    for col in ["B", "C", "D", "E", "F", "G", "H"]:
        ws.column_dimensions[col].width = 18

def create_class_load_sheet(wb, class_loads):
    ws = wb.create_sheet(title="Class Load")
    ws.views.sheetView[0].showGridLines = True
    
    # Title Banner
    ws.merge_cells("A1:J1")
    ws["A1"] = "CLASS LOAD DIRECTORY (1st24)"
    ws["A1"].font = Font(name="Calibri", size=15, bold=True, color=TEXT_LIGHT)
    ws["A1"].fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35
    
    # Table Headers
    headers = [
        "Subject Name", "Section Code", "Room", "Day", "Start Time", 
        "End Time", "Type", "Units", "Hours", "Load Category"
    ]
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=c_idx, value=h)
        cell.font = Font(bold=True, color=TEXT_LIGHT)
        cell.fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        cell.border = get_thin_border()
        
    ws.row_dimensions[3].height = 25
    
    # Data Rows
    thin_border = get_thin_border()
    num_rows = len(class_loads)
    for r_idx, row in class_loads.iterrows():
        row_num = r_idx + 4
        ws.row_dimensions[row_num].height = 20
        
        ws.cell(row=row_num, column=1, value=row["subject_name"]).border = thin_border
        ws.cell(row=row_num, column=2, value=row["section_code"]).border = thin_border
        ws.cell(row=row_num, column=3, value=row["room"]).border = thin_border
        
        day_cell = ws.cell(row=row_num, column=4, value=row["day_of_week"])
        day_cell.alignment = Alignment(horizontal="center")
        day_cell.border = thin_border
        
        start_cell = ws.cell(row=row_num, column=5, value=row["start_time"])
        start_cell.alignment = Alignment(horizontal="center")
        start_cell.border = thin_border
        
        end_cell = ws.cell(row=row_num, column=6, value=row["end_time"])
        end_cell.alignment = Alignment(horizontal="center")
        end_cell.border = thin_border
        
        type_cell = ws.cell(row=row_num, column=7, value=row["type"])
        type_cell.alignment = Alignment(horizontal="center")
        type_cell.border = thin_border
        
        units_cell = ws.cell(row=row_num, column=8, value=row["units"])
        units_cell.alignment = Alignment(horizontal="right")
        units_cell.border = thin_border
        
        hours_cell = ws.cell(row=row_num, column=9, value=row["hours"])
        hours_cell.alignment = Alignment(horizontal="right")
        hours_cell.border = thin_border
        
        cat_cell = ws.cell(row=row_num, column=10, value=row["load_category"])
        cat_cell.border = thin_border
        
    # Totals Row
    total_row = num_rows + 4
    ws.row_dimensions[total_row].height = 22
    
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
    tot_label = ws.cell(row=total_row, column=1, value="Total")
    tot_label.font = Font(bold=True)
    tot_label.alignment = Alignment(horizontal="right")
    
    # Border for the merged cell
    for col in range(1, 8):
        ws.cell(row=total_row, column=col).border = thin_border
        ws.cell(row=total_row, column=col).fill = PatternFill(start_color=SECONDARY_COLOR, end_color=SECONDARY_COLOR, fill_type="solid")
        
    # Excel Formulas
    # Units formula
    cell_units = ws.cell(row=total_row, column=8, value=f"=SUM(H4:H{total_row-1})")
    cell_units.font = Font(bold=True)
    cell_units.border = thin_border
    cell_units.fill = PatternFill(start_color=SECONDARY_COLOR, end_color=SECONDARY_COLOR, fill_type="solid")
    
    # Hours formula
    cell_hours = ws.cell(row=total_row, column=9, value=f"=SUM(I4:I{total_row-1})")
    cell_hours.font = Font(bold=True)
    cell_hours.border = thin_border
    cell_hours.fill = PatternFill(start_color=SECONDARY_COLOR, end_color=SECONDARY_COLOR, fill_type="solid")
    
    # Category cell empty
    ws.cell(row=total_row, column=10).border = thin_border
    ws.cell(row=total_row, column=10).fill = PatternFill(start_color=SECONDARY_COLOR, end_color=SECONDARY_COLOR, fill_type="solid")
    
    # Columns width
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 8
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 8
    ws.column_dimensions["H"].width = 8
    ws.column_dimensions["I"].width = 8
    ws.column_dimensions["J"].width = 18

def create_timesheet_sheet(wb, sheet_title, instructor, class_loads, year, month, cutoff_type, timesheet_data):
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True
    
    # Generate dates for this cutoff
    dates = generate_cutoff_dates(year, month, cutoff_type)
    num_dates = len(dates)
    last_col_idx = 2 + num_dates
    last_col_letter = get_column_letter(last_col_idx)
    total_col_idx = last_col_idx + 1
    total_col_letter = get_column_letter(total_col_idx)
    
    # Title Block
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_col_idx)
    ws["A1"] = f"ACADEMIC TIMESHEET ({cutoff_type})"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=TEXT_LIGHT)
    ws["A1"].fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40
    
    # Header metadata
    month_name = calendar.month_name[month]
    acad_level = instructor.get("academic_level", "Tertiary")
    headers_meta = [
        ("Instructor Name:", instructor["name"], "Employment Status:", f"{instructor['employment_status']} ({acad_level})"),
        ("Cutoff Period:", f"{month_name} {cutoff_type}, {year}", "Department:", instructor.get("department", "N/A"))
    ]
    for r_idx, row_data in enumerate(headers_meta, start=2):
        ws.row_dimensions[r_idx].height = 20
        ws.cell(row=r_idx, column=1, value=row_data[0]).font = Font(bold=True)
        ws.cell(row=r_idx, column=2, value=row_data[1])
        ws.cell(row=r_idx, column=5, value=row_data[2]).font = Font(bold=True)
        ws.cell(row=r_idx, column=6, value=row_data[3])
        
    # Table headers: Date values and Weekdays
    ws.cell(row=5, column=1, value="Activity Category").font = Font(bold=True, color=TEXT_LIGHT)
    ws.cell(row=5, column=1).fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    ws.cell(row=5, column=1).alignment = Alignment(vertical="center")
    ws.cell(row=5, column=1).border = get_thin_border()
    
    ws.cell(row=5, column=2, value="Activity Details").font = Font(bold=True, color=TEXT_LIGHT)
    ws.cell(row=5, column=2).fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    ws.cell(row=5, column=2).alignment = Alignment(vertical="center")
    ws.cell(row=5, column=2).border = get_thin_border()
    
    ws.cell(row=6, column=1).border = get_thin_border()
    ws.cell(row=6, column=2).border = get_thin_border()
    ws.merge_cells("A5:A6")
    ws.merge_cells("B5:B6")
    
    # Fill Date Columns
    thin_border = get_thin_border()
    for idx, d_info in enumerate(dates):
        c_idx = 3 + idx
        c_let = get_column_letter(c_idx)
        
        # Row 5: Day number
        d_cell = ws.cell(row=5, column=c_idx, value=d_info['day_num'])
        d_cell.font = Font(bold=True)
        d_cell.alignment = Alignment(horizontal="center")
        d_cell.border = thin_border
        
        # Row 6: Day Abbreviation
        abbr_cell = ws.cell(row=6, column=c_idx, value=d_info['day_abbr'])
        abbr_cell.font = Font(bold=True, size=9)
        abbr_cell.alignment = Alignment(horizontal="center")
        abbr_cell.border = thin_border
        
        # Highlight Saturdays & Sundays
        if d_info['day_abbr'] in ("S", "SU"):
            d_cell.fill = PatternFill(start_color=WEEKEND_COLOR, end_color=WEEKEND_COLOR, fill_type="solid")
            abbr_cell.fill = PatternFill(start_color=WEEKEND_COLOR, end_color=WEEKEND_COLOR, fill_type="solid")
            
    # Row 5/6 Total Column
    ws.merge_cells(start_row=5, start_column=total_col_idx, end_row=6, end_column=total_col_idx)
    tot_header = ws.cell(row=5, column=total_col_idx, value="Total Hours")
    tot_header.font = Font(bold=True, color=TEXT_LIGHT)
    tot_header.fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    tot_header.alignment = Alignment(horizontal="center", vertical="center")
    tot_header.border = thin_border
    
    ws.row_dimensions[5].height = 18
    ws.row_dimensions[6].height = 18
    
    # Construct rows of activities
    # Separate Regular Load and Excess Load
    regular_classes = class_loads[class_loads["load_category"] == "Regular Load"]
    excess_classes = class_loads[class_loads["load_category"].isin(["Excess Load", "Excess/Overload"])]
    
    activity_rows = []
    
    # Section A: Regular Load
    for _, cl in regular_classes.iterrows():
        key_name = f"{cl['type']}: {cl['subject_name']} ({cl['section_code']})"
        activity_rows.append(("Regular Load", key_name, key_name))
        
    # Section B: Excess Load
    for _, cl in excess_classes.iterrows():
        key_name = f"{cl['type']}: {cl['subject_name']} ({cl['section_code']})"
        activity_rows.append(("Excess Load", key_name, key_name))
        
    # Section C: Non-Teaching SC
    activity_rows.append(("Non-Teaching SC", "Career Orientation Seminars (COS)", "Career Orientation Seminars (COS)"))
    activity_rows.append(("Non-Teaching SC", "Guidance/Counseling", "Guidance/Counseling"))
    
    # Section D: Non-Teaching HQ
    activity_rows.append(("Non-Teaching HQ", "Consultation", "Consultation"))
    activity_rows.append(("Non-Teaching HQ", "Administrative Hours", "Administrative Hours"))
    
    # Write Row Contents
    start_data_row = 7
    for idx, (cat, detail, key_name) in enumerate(activity_rows):
        r_num = start_data_row + idx
        ws.row_dimensions[r_num].height = 20
        
        c_cell = ws.cell(row=r_num, column=1, value=cat)
        c_cell.border = thin_border
        
        d_cell = ws.cell(row=r_num, column=2, value=detail)
        d_cell.border = thin_border
        
        # Write hours for each date
        for d_idx, d_info in enumerate(dates):
            col_idx = 3 + d_idx
            date_str = d_info['date'].strftime("%Y-%m-%d")
            
            # Check for overrides first
            val = timesheet_data.get((key_name, date_str))
            
            if val is None:
                # If no override, auto-populate regular / excess / consultation
                if cat == "Regular Load":
                    val = 0.0
                    for _, cl in regular_classes.iterrows():
                        cl_key = f"{cl['type']}: {cl['subject_name']} ({cl['section_code']})"
                        if cl_key == key_name and cl['day_of_week'] == d_info['day_abbr']:
                            val = cl['hours']
                            break
                elif cat == "Excess Load":
                    val = 0.0
                    for _, cl in excess_classes.iterrows():
                        cl_key = f"{cl['type']}: {cl['subject_name']} ({cl['section_code']})"
                        if cl_key == key_name and cl['day_of_week'] == d_info['day_abbr']:
                            val = cl['hours']
                            break
                elif cat == "Teaching Activities": # Fallback for backwards compatibility
                    val = 0.0
                    for _, cl in pd.concat([regular_classes, excess_classes]).iterrows():
                        cl_key = f"{cl['type']}: {cl['subject_name']} ({cl['section_code']})"
                        if cl_key == key_name and cl['day_of_week'] == d_info['day_abbr']:
                            val = cl['hours']
                            break
                elif key_name == "Consultation":
                    # Pre-populate consultation from scheduled consultation
                    val = 0.0
                    consultations = class_loads[class_loads["load_category"] == "Consultation"]
                    for _, cl in consultations.iterrows():
                        if cl['day_of_week'] == d_info['day_abbr']:
                            val = cl['hours']
                            break
                else:
                    val = 0.0
                    
            # Write value
            h_cell = ws.cell(row=r_num, column=col_idx, value=val)
            h_cell.alignment = Alignment(horizontal="right")
            h_cell.border = thin_border
            
            # Shading for weekend cells in data rows
            if d_info['day_abbr'] in ("S", "SU"):
                h_cell.fill = PatternFill(start_color=WEEKEND_COLOR, end_color=WEEKEND_COLOR, fill_type="solid")
                
        # Total Formula for the row
        tot_formula = f"=SUM(C{r_num}:{last_col_letter}{r_num})"
        t_cell = ws.cell(row=r_num, column=total_col_idx, value=tot_formula)
        t_cell.font = Font(bold=True)
        t_cell.alignment = Alignment(horizontal="right")
        t_cell.border = thin_border
        t_cell.fill = PatternFill(start_color=SECONDARY_COLOR, end_color=SECONDARY_COLOR, fill_type="solid")
        
    # Daily Totals Row
    total_row_idx = start_data_row + len(activity_rows)
    ws.row_dimensions[total_row_idx].height = 22
    
    # Labels
    ws.cell(row=total_row_idx, column=1, value="Daily Totals").font = Font(bold=True)
    ws.cell(row=total_row_idx, column=1).border = thin_border
    ws.cell(row=total_row_idx, column=1).fill = PatternFill(start_color=SECONDARY_COLOR, end_color=SECONDARY_COLOR, fill_type="solid")
    
    ws.cell(row=total_row_idx, column=2, value="").border = thin_border
    ws.cell(row=total_row_idx, column=2).fill = PatternFill(start_color=SECONDARY_COLOR, end_color=SECONDARY_COLOR, fill_type="solid")
    ws.merge_cells(start_row=total_row_idx, start_column=1, end_row=total_row_idx, end_column=2)
    
    # Formulas for columns
    for d_idx, d_info in enumerate(dates):
        c_idx = 3 + d_idx
        col_let = get_column_letter(c_idx)
        daily_formula = f"=SUM({col_let}{start_data_row}:{col_let}{total_row_idx-1})"
        
        d_tot_cell = ws.cell(row=total_row_idx, column=c_idx, value=daily_formula)
        d_tot_cell.font = Font(bold=True)
        d_tot_cell.alignment = Alignment(horizontal="right")
        d_tot_cell.border = thin_border
        d_tot_cell.fill = PatternFill(start_color=SECONDARY_COLOR, end_color=SECONDARY_COLOR, fill_type="solid")
        
    # Grand Total formula
    grand_formula = f"=SUM({total_col_letter}{start_data_row}:{total_col_letter}{total_row_idx-1})"
    g_cell = ws.cell(row=total_row_idx, column=total_col_idx, value=grand_formula)
    g_cell.font = Font(bold=True, size=11, color="000000")
    g_cell.alignment = Alignment(horizontal="right")
    g_cell.border = thin_border
    g_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid") # Soft Green fill for grand total
    
    # Column width formatting
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 28
    for col in range(3, last_col_idx + 1):
        ws.column_dimensions[get_column_letter(col)].width = 6
    ws.column_dimensions[total_col_letter].width = 12
    
    return ws
